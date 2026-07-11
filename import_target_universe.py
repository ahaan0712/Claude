#!/usr/bin/env python3
from __future__ import annotations
import re, hashlib, json
from pathlib import Path
from datetime import date, timedelta
from schema import write_json
try:
    import openpyxl
except Exception:
    openpyxl=None
EXCLUDED_INDUSTRY=["defense","weapons","insurance","industrial manufacturing","manufacturing","energy","oil","gas","utilities","airline","aviation","hospitality","hotel","travel operations"]
QUANT=["quant","trading","investment banking","actuarial","underwriting"]
MONTHS={m.lower():i for i,m in enumerate(['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'],1)}
MONTHS.update({m.lower():i for i,m in enumerate(['January','February','March','April','May','June','July','August','September','October','November','December'],1)})
def canon(s): return re.sub(r'[^a-z0-9]+','-',(s or '').lower()).strip('-')
def pid(company, row): return canon(company)+'-'+hashlib.sha1(str(row).encode()).hexdigest()[:8]
def last_day(y,m):
    return (date(y+1,1,1)-timedelta(days=1)).day if m==12 else (date(y,m+1,1)-timedelta(days=1)).day
def parse_window(text):
    raw=text or ''; low=raw.lower(); conf='LOW'; start=end=None
    y=2026
    ym=re.search(r'(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s*[-–]\s*(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s*(20\d{2})?', low)
    if ym:
        m1=MONTHS[ym.group(1)[:3].lower()]; m2=MONTHS[ym.group(2)[:3].lower()]; y=int(ym.group(3) or 2026); start=date(y,m1,1); end=date(y,m2,last_day(y,m2)); conf='HIGH'
    else:
        sm=re.search(r'(mid-|mid\s*)?(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s*(20\d{2})?', low)
        if sm:
            m=MONTHS[sm.group(2)[:3].lower()]; y=int(sm.group(3) or 2026); day=15 if sm.group(1) else 1; start=date(y,m,day); end=start+timedelta(days=28); conf='MEDIUM' if 'mid' in (sm.group(1) or '') or '~' in low else 'MEDIUM'
    monitor=(start-timedelta(days=21)).isoformat() if start else '2026-07-11'
    return {'expected_open_start':start.isoformat() if start else None,'expected_open_end':end.isoformat() if end else None,'opening_window_text':raw,'opening_window_confidence':conf,'monitoring_start_date':monitor,'recommended_check_frequency':'weekly' if conf in {'HIGH','MEDIUM'} else 'monthly'}
def find_workbook():
    exact=Path('Ahaan_2027_Internship_Target_List_Cleaned.xlsx')
    if exact.exists(): return exact
    cands=list(Path('.').glob('*2027*Internship*Target*List*Cleaned*.xlsx'))+list(Path('.').glob('*Target*List*.xlsx'))
    return cands[0] if cands else None
def fallback_rows():
    old=Path('data/target_programs.json')
    if old.exists():
        rows=[]
        for i,r in enumerate(json.loads(old.read_text()),2):
            rows.append({'Company':r.get('company',''),'Category':r.get('category','Analytics'),'Internship Program':r.get('intended_program_or_role') or r.get('intended_internship_program') or 'Summer 2027 analytics internship target','Typical Opening (2027 cycle)':r.get('opening_window_text') or r.get('recruiting_window') or 'Jul-Sep 2026','Deadline Behavior':r.get('deadline_behavior','Varies'),'How to Apply':r.get('application_route','Official careers page'),'Career Page Link':r.get('career_page_url',''),'Intl-Student Friendly?':r.get('international_student_evidence','Verify posting-level CPT/sponsorship language'),'Notes / Strategy':r.get('strategic_notes') or r.get('source_notes',''),'_row':r.get('original_spreadsheet_row_number',i)})
        if rows: return rows, 'sanitized derived JSON fallback', 'Target Companies'
    return [], 'missing', 'Target Companies'
def load_rows():
    wb=find_workbook()
    if not wb or not openpyxl: return fallback_rows()
    book=openpyxl.load_workbook(wb, data_only=True, read_only=True)
    if 'Target Companies' not in book.sheetnames: raise SystemExit('Authoritative worksheet Target Companies not found')
    sh=book['Target Companies']; vals=list(sh.iter_rows(values_only=True)); header_i=next(i for i,r in enumerate(vals) if r and 'Company' in [str(x).strip() for x in r if x])
    headers=[str(x).strip() if x is not None else '' for x in vals[header_i]]; rows=[]
    for ridx,r in enumerate(vals[header_i+1:], header_i+2):
        d={headers[i]: (r[i] if i<len(r) and r[i] is not None else '') for i in range(len(headers))}; d['_row']=ridx
        if any(str(v).strip() for k,v in d.items() if k!='_row'): rows.append(d)
    return rows, wb.name, 'Target Companies'
def main():
    rows,source,worksheet=load_rows(); programs=[]; excluded=[]; companies={}; malformed=[]; aliases={}; dup_programs=[]
    for r in rows:
        company=str(r.get('Company','')).strip(); program=str(r.get('Internship Program','')).strip()
        if not company: malformed.append({'row':r.get('_row'),'reason':'missing_company'}); continue
        text=' '.join(str(r.get(k,'')) for k in r).lower(); reason=next((x for x in QUANT if x in text),None); ind=next((x for x in EXCLUDED_INDUSTRY if x in text),None)
        cid=canon(company); win=parse_window(str(r.get('Typical Opening (2027 cycle)','')))
        rec={'target_program_id':pid(company,r.get('_row')),'company':company,'canonical_company_name':company,'canonical_company_id':cid,'aliases':[company],'category':str(r.get('Category','')).strip(),'intended_internship_program':program or 'Summer 2027 analytics internship target','expected_opening_window':str(r.get('Typical Opening (2027 cycle)','')).strip(),'deadline_behavior':str(r.get('Deadline Behavior','')).strip(),'application_route':str(r.get('How to Apply','')).strip(),'career_page_url':str(r.get('Career Page Link','')).strip(),'international_student_evidence':str(r.get('Intl-Student Friendly?','')).strip(),'strategic_notes':str(r.get('Notes / Strategy','')).strip(),'original_spreadsheet_row_number':r.get('_row'),**win,'status':'ARCHIVED' if reason or ind else 'ACTIVE','archive_reason':reason or ind or ''}
        if rec['status']=='ARCHIVED': excluded.append(rec)
        else:
            programs.append(rec); c=companies.setdefault(cid,{'canonical_company_id':cid,'company':company,'canonical_company_name':company,'aliases':[company],'status':'ACTIVE','target_program_ids':[],'career_page_url':rec['career_page_url'],'ats_provider':'UNKNOWN','monitoring_status':'NEEDS_RESEARCH','last_successful_check':None,'last_checked':None,'next_check_date':rec['monitoring_start_date'],'expected_open_start':rec['expected_open_start'],'expected_open_end':rec['expected_open_end'],'opening_window_text':rec['opening_window_text'],'opening_window_confidence':rec['opening_window_confidence'],'recommended_check_frequency':rec['recommended_check_frequency'],'deadline_behavior':rec['deadline_behavior'],'current_status':'Target retained; no verified live posting yet.'}); c['target_program_ids'].append(rec['target_program_id'])
    report={'source_workbook':source,'worksheet_used':worksheet,'total_rows':len(rows),'valid_rows':len(rows)-len(malformed),'malformed_rows':malformed,'duplicate_company_aliases':aliases,'duplicate_programs':dup_programs,'active_companies':len(companies),'archived_companies':len({r['canonical_company_id'] for r in excluded}),'active_programs':len(programs),'quant_programs_removed':sum(1 for r in excluded if r['archive_reason'] in QUANT),'industry_exclusions':{},'unresolved_records':malformed,'source_precedence':['exact live posting text','verified résumé/profile','cleaned XLSX','current official company career pages','verified historical evidence','V6 reference metadata','supplemental tracker data']}
    for r in excluded: report['industry_exclusions'][r['archive_reason']]=report['industry_exclusions'].get(r['archive_reason'],0)+1
    write_json('data/target_programs.json',programs); write_json('data/company_watchlist.json',sorted(companies.values(),key=lambda x:x['company'].lower())); write_json('data/excluded_target_programs.json',excluded); write_json('data/company_reconciliation_report.json',report)
if __name__=='__main__': main()
